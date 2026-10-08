from io import BytesIO
from tempfile import TemporaryDirectory
from unittest.mock import Mock, patch

from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import SimpleTestCase, override_settings
from django.urls import reverse
from openpyxl import Workbook, load_workbook
from PIL import Image
from pypdf import PdfReader
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle
from rest_framework.test import APITestCase

from api.models import Brand, Product
from .ai_parsing import AI_PARSE_JSON_SCHEMA, _bind_result_source, _normalize_ai_result
from .excel import build_quotation_excel
from .import_parsers import parse_file_preview, parse_text_preview
from .models import AIParseCache, Company, InquiryLine, Quotation, QuotationSettings
from .pdf import build_quotation_pdf


HEADERS = ["No", "Description", "Brand", "Qty", "Unit", "Expiry", "Price", "VAT"]
ROWS = [
    [1, "Adhesive tape", "3M", 2, "box", "09/2028", 10, 5],
    [2, "Gauze swabs", "Acme Medical", 5, "pack", "", 12, 0],
    [3, "Eye wash solution", "", 1, "each", "03/2029", 8, 0],
]


def brand_upload(header="Brand"):
    book = Workbook()
    book.active.append([*HEADERS[:2], header, *HEADERS[3:]])
    for row in ROWS:
        book.active.append(row)
    stream = BytesIO()
    book.save(stream)
    return SimpleUploadedFile("brands.xlsx", stream.getvalue())


class InquiryBrandParserTests(SimpleTestCase):
    def assert_rows(self, preview):
        self.assertEqual([r.get("brand_name") for r in preview["lines"]], ["3M", "Acme Medical", ""])
        self.assertEqual([r["raw_name"].lower() for r in preview["lines"]], [r[1].lower() for r in ROWS])
        self.assertEqual([r["quantity"] for r in preview["lines"]], ["2", "5", "1"])
        self.assertEqual([r["unit_price"] for r in preview["lines"]], ["10", "12", "8"])
        self.assertEqual(preview["lines"][0]["expiry_date"], "09/2028")

    def test_excel_brand_header_aliases(self):
        for header in ("Brand", "Brand name", "Product Brand", "Make", "Brand / Make"):
            with self.subTest(header=header):
                self.assert_rows(parse_file_preview(brand_upload(header), store_source=False))

    def test_pasted_table_and_html_preserve_brand_cells(self):
        text = "\n".join("\t".join(map(str, row)) for row in [HEADERS, *ROWS])
        self.assert_rows(parse_text_preview(text))
        html = "<table>" + "".join("<tr>" + "".join(f"<td>{cell}</td>" for cell in row) + "</tr>" for row in [HEADERS, *ROWS]) + "</table>"
        self.assert_rows(parse_text_preview(text, raw_html=html))

    def test_pdf_table_preserves_brand_and_expiry(self):
        stream = BytesIO()
        table = Table([HEADERS, *ROWS], colWidths=[25, 110, 95, 25, 35, 65, 40, 30])
        table.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.5, colors.black)]))
        SimpleDocTemplate(stream).build([table])
        self.assert_rows(parse_file_preview(SimpleUploadedFile("brands.pdf", stream.getvalue()), store_source=False))

    def test_no_brand_column_does_not_treat_description_or_supplier_as_brand(self):
        preview = parse_text_preview("Supplier: Acme Medical\nItem\tQty\tUnit\n3M adhesive tape\t2\tbox")
        self.assertEqual(len(preview["lines"]), 1)
        self.assertFalse(preview["lines"][0].get("brand_name"))
        self.assertIn("3M", preview["lines"][0]["raw_name"])

    def test_ai_text_and_vision_keep_brand_separate_and_leave_absent_brand_blank(self):
        self.assertIn("brand_name", AI_PARSE_JSON_SCHEMA["properties"]["rows"]["items"]["required"])
        for mode in (AIParseCache.MODE_TEXT, AIParseCache.MODE_VISION):
            with self.subTest(mode=mode):
                result = _normalize_ai_result({"rows": [
                    {"item_name": "Adhesive tape", "brand_name": "3M", "quantity": "2", "unit": "box", "parse_status": "parsed", "confidence": 99},
                    {"item_name": "Eye wash solution", "quantity": "1", "unit": "each", "parse_status": "parsed", "confidence": 99},
                ]}, preview={}, mode=mode, provider="test", model="test", output_style="inquiry")
                self.assertEqual([r["brand_name"] for r in result["lines"]], ["3M", ""])
                self.assertEqual(result["lines"][0]["raw_name"], "Adhesive Tape")

    def test_cleanup_preserves_brands_and_blanks_after_reordering(self):
        preview = parse_file_preview(brand_upload(), store_source=False)
        candidate = {"lines": [dict(row, brand_name="Wrong brand") for row in reversed(preview["lines"])]}
        result = _bind_result_source(candidate, preview)
        self.assertEqual([r["brand_name"] for r in result["lines"]], ["", "Acme Medical", "3M"])
        self.assertEqual(candidate["lines"][0]["brand_name"], "Wrong brand")
        # Isolate brand provenance so the brand guard, rather than expiry, rejects missing rows.
        for row in preview["lines"]:
            row.pop("expiry_date", None)
        result = _bind_result_source({"lines": candidate["lines"][:1]}, preview)
        self.assertEqual(result["lines"], preview["lines"])
        self.assertEqual(result["meta"]["ai_cleanup_rejection_reason"], "brand_evidence_unmapped")


@override_settings(QUOTATION_AI_PARSE_GLOBAL_ENABLED=False)
class InquiryBrandWorkflowTests(APITestCase):
    def setUp(self):
        self.staff = get_user_model().objects.create_user(username="brand_import_staff", is_staff=True)
        self.client.force_authenticate(self.staff)
        self.company = Company.objects.create(name="Brand importer")
        temp = TemporaryDirectory()
        self.addCleanup(temp.cleanup)
        settings = override_settings(QUOTATION_PRIVATE_STORAGE_ROOT=temp.name)
        settings.enable()
        self.addCleanup(settings.disable)

    def test_upload_review_save_reopen_and_exports_keep_brand(self):
        parsed = self.client.post(reverse("quotation-inquiry-parse-file"), {"file": brand_upload()}, format="multipart")
        self.assertEqual(parsed.status_code, 200, parsed.data)
        lines = parsed.data["lines"]
        self.assertEqual([r["brand_name"] for r in lines], ["3M", "Acme Medical", ""])
        # Mirror the review screen's submission of detected source pricing.
        for line in lines:
            line["unit_price"] = line.get("customer_unit_price") or None
            line["vat_rate"] = line.get("customer_vat_rate") or "0"
        # A reviewed source brand belongs to this document, not the catalogue.
        catalogue_brand = Brand.objects.create(name="Catalogue Brand")
        product = Product.objects.create(name="Adhesive tape", price="10", brand=catalogue_brand, status="draft")
        lines[0].update(brand_name="Reviewed Brand", matched_product=product.pk)
        created = self.client.post(reverse("quotation-inquiry-create-imported"), {"company": self.company.pk, "source_type": "excel", "lines": lines}, format="json")
        self.assertEqual(created.status_code, 201, created.data)
        self.assertEqual(created.data["lines"][0]["brand_name"], "Reviewed Brand")
        self.assertEqual(InquiryLine.objects.get(pk=created.data["lines"][0]["id"]).brand_name, "Reviewed Brand")
        response = self.client.post(reverse("quotation-inquiry-create-quote", args=[created.data["id"]]))
        self.assertEqual(response.status_code, 201, response.data)
        reopened = self.client.get(reverse("quotation-detail", args=[response.data["id"]])).data
        self.assertTrue(reopened["show_brand_column"])
        self.assertTrue(reopened["show_expiry_column"])
        self.assertEqual([r["brand_name_snapshot"] for r in reopened["lines"]], ["Reviewed Brand", "Acme Medical", ""])
        product.refresh_from_db()
        self.assertEqual(product.brand_id, catalogue_brand.pk)
        quote = Quotation.objects.get(pk=response.data["id"])
        pdf = build_quotation_pdf(quote)
        pdf_text = " ".join(page.extract_text() for page in PdfReader(BytesIO(pdf)).pages)
        self.assertIn("Reviewed Brand", " ".join(pdf_text.split()))
        book = load_workbook(BytesIO(build_quotation_excel(quote)))
        self.assertIn("Reviewed Brand", [cell.value for row in book.active for cell in row])

    @override_settings(
        QUOTATION_AI_PARSE_GLOBAL_ENABLED=True,
        QUOTATION_AI_PARSE_PROVIDER="openai",
        QUOTATION_AI_PARSE_VISION_MODEL="test-vision-model",
    )
    @patch.dict("os.environ", {"OPENAI_API_KEY": "test-key"}, clear=False)
    def test_vision_upload_brand_survives_review_and_quote_creation(self):
        settings = QuotationSettings.get_solo()
        settings.ai_parsing_enabled = True
        settings.ai_pdf_vision_enabled = True
        settings.save()
        image = BytesIO()
        Image.new("RGB", (16, 16), "white").save(image, format="PNG")
        provider = Mock()
        provider.clean_rows.return_value = ({"rows": [{
            "item_name": "Adhesive tape", "brand_name": "3M", "quantity": "2",
            "unit": "box", "parse_status": "parsed", "confidence": 0.99,
            "raw_source_text": "Adhesive tape | 3M | 2 | box",
        }]}, {"input_tokens": 10, "output_tokens": 15})
        with patch("quotations.ai_parsing.get_ai_parse_provider", return_value=provider):
            parsed = self.client.post(reverse("quotation-inquiry-parse-file"), {
                "file": SimpleUploadedFile("brands.png", image.getvalue(), content_type="image/png"),
            }, format="multipart")
        self.assertEqual(parsed.status_code, 200, parsed.data)
        self.assertEqual(parsed.data["result_source"], "ai_vision_cleanup")
        self.assertEqual(parsed.data["lines"][0]["brand_name"], "3M")
        lines = [dict(row, vat_rate=row.get("vat_rate") or "0") for row in parsed.data["lines"]]
        created = self.client.post(reverse("quotation-inquiry-create-imported"), {
            "company": self.company.pk, "source_type": "image", "lines": lines,
        }, format="json")
        self.assertEqual(created.status_code, 201, created.data)
        quote = self.client.post(reverse("quotation-inquiry-create-quote", args=[created.data["id"]]))
        self.assertEqual(quote.status_code, 201, quote.data)
        self.assertTrue(quote.data["show_brand_column"])
        self.assertEqual(quote.data["lines"][0]["brand_name_snapshot"], "3M")

    def test_absent_cleared_and_ignored_brands_do_not_enable_column(self):
        for extra in ({}, {"brand_name": ""}, {"brand_name": "3M", "match_status": "ignored"}):
            with self.subTest(extra=extra):
                created = self.client.post(reverse("quotation-inquiry-create-imported"), {
                    "company": self.company.pk, "source_type": "pasted_text",
                    "lines": [{"raw_name": "Gauze", "quantity": "1", **extra}],
                }, format="json")
                self.assertEqual(created.status_code, 201, created.data)
                quote = self.client.post(reverse("quotation-inquiry-create-quote", args=[created.data["id"]]))
                self.assertFalse(quote.data["show_brand_column"])

    def test_overlong_brand_is_rejected(self):
        response = self.client.post(reverse("quotation-inquiry-create-imported"), {
            "company": self.company.pk, "source_type": "pasted_text",
            "lines": [{"raw_name": "Gauze", "quantity": "1", "brand_name": "x" * 201}],
        }, format="json")
        self.assertEqual(response.status_code, 400)
        self.assertFalse(InquiryLine.objects.exists())
