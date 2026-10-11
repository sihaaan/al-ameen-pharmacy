from dataclasses import replace
from datetime import date, datetime, timezone
from io import BytesIO
from unittest.mock import patch

from django.test import TestCase, override_settings
from openpyxl import load_workbook
from pypdf import PdfReader

from .excel import build_quotation_excel
from .models import Company, Quotation, QuotationLine
from .pdf import build_quotation_pdf
from .pdf_config import get_quotation_pdf_config


@override_settings(TIME_ZONE="Asia/Dubai")
class QuotationValidityTests(TestCase):
    def setUp(self):
        self.quote = Quotation.objects.create(company=Company.objects.create(name="Validity Customer"))
        # The quotation date is 11 October in Dubai, despite the UTC date being the 10th.
        self.quote.created_at = datetime(2026, 10, 10, 22, tzinfo=timezone.utc)
        QuotationLine.objects.create(
            quotation=self.quote, item_name_snapshot="Examination gloves", quantity=1,
            unit="box", unit_price=10, vat_rate=5,
        )

    def assert_validity(self, selected_date, default_days, expected_date, expected_footer):
        self.quote.valid_until = selected_date
        config = replace(get_quotation_pdf_config(), validity_days=default_days)
        pdf = PdfReader(BytesIO(build_quotation_pdf(self.quote, config=config)))
        text = " ".join(page.extract_text() for page in pdf.pages)
        self.assertRegex(text, rf"Valid Until\s+{expected_date.isoformat()}")
        self.assertIn(f"Validity: Valid until {expected_footer}.", text)
        self.assertNotIn("days from quotation date unless otherwise stated", text)
        with patch("quotations.excel.get_quotation_pdf_config", return_value=config):
            sheet = load_workbook(BytesIO(build_quotation_excel(self.quote)))["Quotation"]
        row = next(row for row in sheet.iter_rows() if row[0].value == "Valid Until")
        self.assertEqual(row[1].value.date(), expected_date)

    def test_selected_validity_date_overrides_default_in_header_footer_and_excel(self):
        self.assert_validity(date(2026, 12, 31), 30, date(2026, 12, 31), "31 December 2026")

    def test_fresh_model_with_iso_date_text_has_matching_pdf_validity(self):
        quote = Quotation.objects.create(company=self.quote.company, valid_until="2026-12-31")
        text = " ".join(page.extract_text() for page in PdfReader(BytesIO(build_quotation_pdf(quote))).pages)
        self.assertRegex(text, r"Valid Until\s+2026-12-31")
        self.assertIn("Validity: Valid until 31 December 2026.", text)

    def test_unspecified_validity_uses_default_days_from_local_quotation_date(self):
        self.assert_validity(None, 30, date(2026, 11, 10), "10 November 2026")

    def test_custom_default_validity_is_consistent(self):
        self.assert_validity(None, 14, date(2026, 10, 25), "25 October 2026")
