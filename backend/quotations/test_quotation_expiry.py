from io import BytesIO

from django.contrib.auth import get_user_model
from django.urls import reverse
from openpyxl import load_workbook
from pypdf import PdfReader
from reportlab.lib.units import mm
from rest_framework.test import APITestCase

from api.models import Product
from .excel import build_quotation_excel
from .models import Company, Quotation, QuotationLine
from .pdf import build_quotation_pdf, build_proforma_invoice_pdf, _quotation_line_column_widths
from .quotation_email_delivery import quotation_review_fingerprint
from .services import build_quotation_delete_snapshot, recalculate_quotation_totals


class QuotationExpiryTests(APITestCase):
    def setUp(self):
        self.staff = get_user_model().objects.create_user(username="expiry_staff", is_staff=True)
        self.client.force_authenticate(self.staff)
        self.quote = Quotation.objects.create(company=Company.objects.create(name="Expiry Customer"), created_by=self.staff)
        self.product = Product.objects.create(name="Gauze", price="10", status="draft")
        self.line = QuotationLine.objects.create(
            quotation=self.quote, product=self.product, item_name_snapshot="Gauze 7.5 cm",
            brand_name_snapshot="Medical Brand", quantity="2", unit="box", unit_price="10",
            vat_rate="5", match_status=QuotationLine.MATCH_CONFIRMED,
        )
        recalculate_quotation_totals(self.quote)

    def test_toggle_line_save_hide_and_revision_preserve_expiry(self):
        detail = reverse("quotation-detail", args=[self.quote.pk])
        self.assertFalse(self.client.get(detail).data["show_expiry_column"])
        self.assertEqual(self.client.patch(detail, {"show_expiry_column": True}, format="json").status_code, 200)
        created = self.client.post(reverse("quotation-line-list"), {
            "quotation": self.quote.pk, "product": self.product.pk, "item_name_snapshot": "Gloves",
            "quantity": "1", "unit_price": "5", "vat_rate": "0", "match_status": "confirmed",
            "expiry_date": "09/2028",
        }, format="json")
        self.assertEqual(created.status_code, 201, created.data)
        self.assertEqual(created.data["expiry_date"], "09/2028")
        saved = self.client.post(reverse("quotation-bulk-update-lines", args=[self.quote.pk]), {
            "lines": [{"id": self.line.pk, "expiry_date": " 30/09/2028 "}],
        }, format="json")
        self.assertEqual(saved.status_code, 200, saved.data)
        self.client.patch(detail, {"show_expiry_column": False}, format="json")
        hidden = self.client.get(detail).data
        self.assertFalse(hidden["show_expiry_column"])
        self.assertEqual(hidden["lines"][0]["expiry_date"], "30/09/2028")
        self.client.patch(detail, {"show_expiry_column": True}, format="json")
        self.assertEqual(self.client.post(reverse("quotation-finalize", args=[self.quote.pk])).status_code, 200)
        locked = self.client.patch(reverse("quotation-line-detail", args=[self.line.pk]), {"expiry_date": "01/2029"}, format="json")
        self.assertEqual(locked.status_code, 400)
        revision = self.client.post(reverse("quotation-revise", args=[self.quote.pk]))
        self.assertEqual(revision.status_code, 201, revision.data)
        self.assertTrue(revision.data["show_expiry_column"])
        self.assertEqual([line["expiry_date"] for line in revision.data["lines"]], ["30/09/2028", "09/2028"])

    def test_expiry_validation_is_atomic_and_changes_invalidate_review(self):
        before = quotation_review_fingerprint(self.quote)
        saved = self.client.post(reverse("quotation-bulk-update-lines", args=[self.quote.pk]), {
            "lines": [{"id": self.line.pk, "expiry_date": "09/2028"}],
        }, format="json")
        self.assertEqual(saved.status_code, 200, saved.data)
        self.quote.refresh_from_db()
        after_line = quotation_review_fingerprint(self.quote)
        self.assertNotEqual(before, after_line)
        self.quote.show_expiry_column = True
        self.quote.save()
        self.assertNotEqual(after_line, quotation_review_fingerprint(self.quote))
        invalid = self.client.post(reverse("quotation-bulk-update-lines", args=[self.quote.pk]), {
            "lines": [{"id": self.line.pk, "quantity": "3", "expiry_date": "x" * 41}],
        }, format="json")
        self.assertEqual(invalid.status_code, 400)
        self.line.refresh_from_db()
        self.assertEqual(self.line.quantity, 2)
        self.assertEqual(self.line.expiry_date, "09/2028")
        invalid_single = self.client.patch(reverse("quotation-line-detail", args=[self.line.pk]), {"expiry_date": "x" * 41}, format="json")
        self.assertEqual(invalid_single.status_code, 400)
        snapshot = build_quotation_delete_snapshot(self.quote)
        self.assertTrue(snapshot["quotation"]["show_expiry_column"])
        self.assertEqual(snapshot["lines"][0]["expiry_date"], "09/2028")

    def test_pdf_and_excel_optional_combinations_keep_totals_and_hide_on_proforma(self):
        self.line.expiry_date = "30/09/2028"
        self.line.save()
        for brand in (False, True):
            for expiry in (False, True):
                with self.subTest(brand=brand, expiry=expiry):
                    self.quote.show_brand_column = brand
                    self.quote.show_expiry_column = expiry
                    self.quote.save()
                    pdf = PdfReader(BytesIO(build_quotation_pdf(self.quote)))
                    text = " ".join(page.extract_text() for page in pdf.pages)
                    self.assertEqual("Expiry date" in text, expiry)
                    self.assertEqual("30/09/2028" in text, expiry)
                    self.assertEqual("Medical Brand" in text, brand)
                    if expiry:
                        self.assertLessEqual(sum(_quotation_line_column_widths(brand, expiry)), 178 * mm + .001)
                    sheet = load_workbook(BytesIO(build_quotation_excel(self.quote)))["Quotation"]
                    headers = [cell.value for cell in sheet[18]]
                    self.assertEqual("Expiry date" in headers, expiry)
                    self.assertEqual(sheet.cell(19, headers.index("Qty") + 1).value, 2)
                    price_cell = sheet.cell(19, headers.index("Unit Price") + 1)
                    self.assertEqual(price_cell.value, 10)
                    self.assertEqual(price_cell.number_format, "#,##0.00#")
                    total_column = headers.index("Line Total") + 1
                    self.assertEqual(sheet.cell(19, total_column).value, 21)
                    total_row = next(row for row in sheet.iter_rows() if row[total_column - 2].value == "Grand Total")
                    self.assertEqual(total_row[total_column - 1].value, 21)
                    if expiry:
                        self.assertEqual(sheet.cell(19, headers.index("Expiry date") + 1).value, "30/09/2028")
        proforma = PdfReader(BytesIO(build_proforma_invoice_pdf(self.quote)))
        self.assertNotIn("30/09/2028", " ".join(page.extract_text() for page in proforma.pages))

    def test_excel_expiry_is_text_and_pdf_repeats_header_on_long_quotes(self):
        self.quote.show_brand_column = self.quote.show_expiry_column = True
        self.quote.save()
        self.line.expiry_date = "=1+1"
        self.line.save()
        sheet = load_workbook(BytesIO(build_quotation_excel(self.quote)))["Quotation"]
        self.assertEqual(sheet["D19"].data_type, "s")
        self.assertEqual(sheet["D19"].value, "=1+1")
        self.line.expiry_date = "09/2028"
        self.line.save()
        for index in range(45):
            QuotationLine.objects.create(
                quotation=self.quote, item_name_snapshot=f"Long product description with size and pack information {index}",
                brand_name_snapshot="Medical supplies international", expiry_date="30/09/2028" if index % 2 else "",
                quantity=1, unit="box", unit_price=10, sort_order=index + 1,
            )
        pdf = PdfReader(BytesIO(build_quotation_pdf(self.quote)))
        self.assertGreater(len(pdf.pages), 1)
        for page in pdf.pages:
            text = page.extract_text()
            if "Long product description" in text:
                self.assertIn("Expiry date", text)
