from io import BytesIO

from django.contrib.auth import get_user_model
from django.urls import reverse
from openpyxl import load_workbook
from pypdf import PdfReader
from rest_framework.test import APITestCase

from api.models import Product
from .excel import build_quotation_excel
from .models import Company, Quotation, QuotationLine
from .pdf import build_quotation_pdf
from .quotation_email_delivery import quotation_review_fingerprint
from .services import build_quotation_delete_snapshot, recalculate_quotation_totals


class QuotationDeliveryTimeTests(APITestCase):
    def setUp(self):
        self.staff = get_user_model().objects.create_user(username="lead_time_staff", is_staff=True)
        self.client.force_authenticate(self.staff)
        self.company = Company.objects.create(name="Lead Time Customer")
        self.quote = Quotation.objects.create(company=self.company, created_by=self.staff)
        self.detail = reverse("quotation-detail", args=[self.quote.pk])
        product = Product.objects.create(name="Gauze", price="10", status="draft")
        QuotationLine.objects.create(
            quotation=self.quote, product=product, item_name_snapshot="Gauze 7.5 cm",
            quantity="2", unit="box", unit_price="10", vat_rate="5",
            match_status=QuotationLine.MATCH_CONFIRMED,
        )
        recalculate_quotation_totals(self.quote)

    def test_new_quote_defaults_to_immediate_and_saves_preset_custom_or_blank(self):
        created = self.client.post(reverse("quotation-list"), {"company": self.company.pk}, format="json")
        self.assertEqual(created.status_code, 201, created.data)
        self.assertEqual(created.data["delivery_time"], "Immediate (subject to stock availability)")
        for lead_time in ("As agreed", "1-2 weeks", "3-6 weeks", "4-8 weeks", "6-10 weeks",
                          "2-3 working days after order confirmation", ""):
            with self.subTest(lead_time=lead_time):
                result = self.client.patch(self.detail, {"delivery_time": lead_time}, format="json")
                self.assertEqual(result.status_code, 200, result.data)
                self.assertEqual(self.client.get(self.detail).data["delivery_time"], lead_time)

    def test_terms_validation_and_stale_review_protect_saved_lead_time(self):
        before = quotation_review_fingerprint(self.quote)
        result = self.client.patch(self.detail, {
            "delivery_time": " 3-6 weeks ", "quotation_review_fingerprint": before,
        }, format="json")
        self.assertEqual(result.status_code, 200, result.data)
        self.quote.refresh_from_db()
        self.assertEqual(self.quote.delivery_time, "3-6 weeks")
        self.assertNotEqual(before, quotation_review_fingerprint(self.quote))
        stale = self.client.patch(self.detail, {
            "delivery_time": "As agreed", "quotation_review_fingerprint": before,
        }, format="json")
        self.assertEqual(stale.status_code, 409, stale.data)
        invalid = self.client.patch(self.detail, {
            "delivery_time": "x" * 201, "show_brand_column": True,
        }, format="json")
        self.assertEqual(invalid.status_code, 400, invalid.data)
        self.quote.refresh_from_db()
        self.assertEqual(self.quote.delivery_time, "3-6 weeks")
        self.assertFalse(self.quote.show_brand_column)
        self.assertEqual(build_quotation_delete_snapshot(self.quote)["quotation"]["delivery_time"], "3-6 weeks")

    def test_finalized_terms_are_locked_and_revision_preserves_them(self):
        self.client.patch(self.detail, {"delivery_time": "4-8 weeks"}, format="json")
        self.assertEqual(self.client.post(reverse("quotation-finalize", args=[self.quote.pk])).status_code, 200)
        locked = self.client.patch(self.detail, {"delivery_time": "As agreed"}, format="json")
        self.assertEqual(locked.status_code, 400, locked.data)
        revision = self.client.post(reverse("quotation-revise", args=[self.quote.pk]))
        self.assertEqual(revision.status_code, 201, revision.data)
        self.assertEqual(revision.data["delivery_time"], "4-8 weeks")
        revised = self.client.patch(reverse("quotation-detail", args=[revision.data["id"]]), {
            "delivery_time": "As agreed",
        }, format="json")
        self.assertEqual(revised.status_code, 200, revised.data)
        self.quote.refresh_from_db()
        self.assertEqual(self.quote.delivery_time, "4-8 weeks")

    def test_exports_include_saved_lead_time_and_omit_unspecified_terms(self):
        for lead_time in ("Immediate (subject to stock availability)", "3-6 weeks", "After <approval> & payment", ""):
            with self.subTest(lead_time=lead_time):
                self.quote.delivery_time = lead_time
                pdf = PdfReader(BytesIO(build_quotation_pdf(self.quote)))
                text = " ".join(page.extract_text() for page in pdf.pages)
                self.assertEqual("Supply lead time:" in text, bool(lead_time))
                if lead_time:
                    self.assertIn(lead_time, text)
                sheet = load_workbook(BytesIO(build_quotation_excel(self.quote)))["Quotation"]
                self.assertEqual(len(sheet.merged_cells.ranges), 0)
                rows = [row for row in sheet.iter_rows() if row[0].value == "Supply lead time"]
                self.assertEqual(len(rows), int(bool(lead_time)))
                if rows:
                    self.assertEqual(rows[0][1].value, lead_time)

    def test_custom_excel_lead_time_cannot_become_a_formula(self):
        self.quote.delivery_time = "=1+1"
        sheet = load_workbook(BytesIO(build_quotation_excel(self.quote)))["Quotation"]
        row = next(row for row in sheet.iter_rows() if row[0].value == "Supply lead time")
        self.assertEqual(row[1].value, "=1+1")
        self.assertEqual(row[1].data_type, "s")
